"""Turn an uploaded file into a list of LangChain Documents (one per page/slide/section)."""
from pathlib import Path
from langchain_core.documents import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def load_file(path: Path, original_name: str | None = None) -> list[Document]:
    ext = path.suffix.lower()
    name = original_name or path.name
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}. Allowed: {sorted(SUPPORTED_EXTENSIONS)}")

    if ext == ".pdf":
        pages = _load_pdf(path)
    elif ext == ".docx":
        pages = _load_docx(path)
    elif ext == ".pptx":
        pages = _load_pptx(path)
    else:
        pages = [(None, path.read_text(encoding="utf-8", errors="ignore"))]

    docs = []
    for page, text in pages:
        text = text.strip()
        if text:
            meta = {"filename": name}
            if page is not None:
                meta["page"] = page
            docs.append(Document(page_content=text, metadata=meta))
    return docs


def _load_pdf(path: Path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]


def _load_docx(path: Path):
    import docx

    d = docx.Document(str(path))
    text = "\n".join(p.text for p in d.paragraphs)
    for table in d.tables:
        for row in table.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)
    return [(None, text)]


def _load_pptx(path: Path):
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        out.append((i, "\n".join(parts)))
    return out
