from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def _style_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")


def _autowidth(ws):
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)


def build_report(req) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Student", "Email", "Score", "Total", "Percentage", "Time (min)", "Submitted At"])
    _style_header(ws)

    pcts = []
    for r in sorted(req.results, key=lambda x: x.student_name.lower()):
        pct = round(r.score / r.total * 100, 1) if r.total else 0
        pcts.append(pct)
        mins = round(r.time_taken_seconds / 60, 1) if r.time_taken_seconds is not None else ""
        ws.append([r.student_name, r.student_email, r.score, r.total, pct, mins, r.submitted_at])
    _autowidth(ws)

    s = wb.create_sheet("Summary")
    s.append(["Quiz", req.quiz_title])
    s.append(["Class", req.class_name])
    s.append(["Students attempted", len(pcts)])
    if pcts:
        s.append(["Average %", round(sum(pcts) / len(pcts), 1)])
        s.append(["Highest %", max(pcts)])
        s.append(["Lowest %", min(pcts)])
        s.append(["Passed (>=40%)", sum(1 for p in pcts if p >= 40)])
    for row in s.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    _autowidth(s)

    if req.question_stats:
        q = wb.create_sheet("Question Analysis")
        q.append(["Question", "Correct", "Attempts", "% Correct"])
        _style_header(q)
        for st in req.question_stats:
            total = st.get("total_attempts", 0)
            correct = st.get("correct_count", 0)
            q.append([st.get("question", ""), correct, total, round(correct / total * 100, 1) if total else 0])
        _autowidth(q)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
