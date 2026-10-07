"""Extract text from PDF / DOCX / PPTX / TXT, page by page."""
import re
from collections import Counter
from dataclasses import dataclass
from io import BytesIO

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation


class ExtractionError(Exception):
    pass


@dataclass
class Page:
    number: int
    text: str


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"-\n(\w)", r"\1", text)       # join hyphenated line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _remove_headers_footers(pages):
    """Drop lines repeated on more than half the pages (page numbers, running titles)."""
    if len(pages) < 4:
        return pages
    counts = Counter()
    for p in pages:
        for line in {l.strip() for l in p.text.split("\n") if l.strip()}:
            counts[re.sub(r"\d+", "#", line)] += 1
    repeated = {k for k, v in counts.items() if v > len(pages) * 0.5}
    for p in pages:
        kept = [l for l in p.text.split("\n") if re.sub(r"\d+", "#", l.strip()) not in repeated]
        p.text = "\n".join(kept)
    return pages


def extract_pdf(data: bytes):
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise ExtractionError("Could not open PDF (corrupted file).")
    if doc.needs_pass:
        raise ExtractionError("PDF is password-protected.")
    pages = [Page(i + 1, _clean(pg.get_text("text"))) for i, pg in enumerate(doc)]
    total_chars = sum(len(p.text) for p in pages)
    if total_chars < 50 * max(1, len(pages)):
        raise ExtractionError("This PDF looks scanned (images, no selectable text). OCR is not enabled.")
    return _remove_headers_footers(pages)


def extract_docx(data: bytes):
    try:
        doc = Document(BytesIO(data))
    except Exception:
        raise ExtractionError("Could not open DOCX file.")
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [Page(1, _clean("\n".join(parts)))]


def extract_pptx(data: bytes):
    try:
        prs = Presentation(BytesIO(data))
    except Exception:
        raise ExtractionError("Could not open PPTX file.")
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
        pages.append(Page(i, _clean("\n".join(texts))))
    return pages


def extract_txt(data: bytes):
    return [Page(1, _clean(data.decode("utf-8", errors="ignore")))]


def extract(filename: str, data: bytes):
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    fn = {"pdf": extract_pdf, "docx": extract_docx, "pptx": extract_pptx,
          "txt": extract_txt, "md": extract_txt}.get(ext)
    if not fn:
        raise ExtractionError(f"Unsupported file type: .{ext}")
    pages = fn(data)
    if sum(len(p.text) for p in pages) < 200:
        raise ExtractionError("Not enough text found in the file.")
    return pages
