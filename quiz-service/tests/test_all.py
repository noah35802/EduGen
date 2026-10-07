from io import BytesIO

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from app.extractor import ExtractionError, extract
from app.generator import chunk_pages, is_valid, pick_chunks, shuffle_options
from app.main import app

client = TestClient(app)
LOREM = ("The network layer is responsible for routing packets between different networks. "
         "Routers use routing tables to decide the next hop. ") * 8


def make_pdf(pages=3, text=LOREM):
    doc = fitz.open()
    for i in range(pages):
        pg = doc.new_page()
        pg.insert_textbox(fitz.Rect(50, 50, 550, 780), f"Page body {i}\n" + " ".join(f"Topic{i}x{j} covers routing concept number {i * 100 + j}." for j in range(40)) + "\n" + text, fontsize=10)
    return doc.tobytes()


def make_docx():
    d = Document()
    d.add_paragraph(LOREM)
    b = BytesIO(); d.save(b); return b.getvalue()


def make_pptx():
    p = Presentation()
    for _ in range(3):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = "Routing"
        s.placeholders[1].text = LOREM
    b = BytesIO(); p.save(b); return b.getvalue()


# ---------- extraction ----------
def test_pdf_pages():
    pages = extract("a.pdf", make_pdf(3))
    assert len(pages) == 3 and "routing" in pages[0].text.lower()

def test_docx_pptx_txt():
    assert extract("a.docx", make_docx())[0].text
    assert len(extract("a.pptx", make_pptx())) == 3
    assert extract("a.txt", LOREM.encode())[0].text

def test_scanned_pdf_rejected():
    doc = fitz.open(); doc.new_page()
    with pytest.raises(ExtractionError):
        extract("scan.pdf", doc.tobytes())

def test_unsupported_and_corrupt():
    with pytest.raises(ExtractionError):
        extract("a.exe", b"x")
    with pytest.raises(ExtractionError):
        extract("a.pdf", b"not a pdf")


# ---------- validation ----------
GOOD = {"question": "Q?", "options": ["a", "b", "c", "d"], "correct_index": 2, "explanation": "because"}

def test_is_valid():
    assert is_valid(GOOD)
    assert not is_valid({**GOOD, "options": ["a", "b", "c"]})
    assert not is_valid({**GOOD, "options": ["a", "a", "c", "d"]})
    assert not is_valid({**GOOD, "correct_index": 4})
    assert not is_valid({**GOOD, "correct_index": True})
    assert not is_valid({**GOOD, "options": ["a", "b", "c", "All of the above"]})
    assert not is_valid({**GOOD, "explanation": ""})

def test_shuffle_keeps_correct_answer():
    for _ in range(50):
        q = shuffle_options({**GOOD, "options": list(GOOD["options"])})
        assert q["options"][q["correct_index"]] == "c"

def test_chunking_and_cap():
    from app.extractor import Page
    pages = [Page(i, f"Unique content for page {i}. " * 60) for i in range(1, 41)]
    chunks = chunk_pages("m1", "a.pdf", pages, size=2000)
    assert len(chunks) > 12
    assert chunks[0].page_start == 1 and chunks[-1].page_end == 40
    assert len(pick_chunks(chunks, 12)) == 12


# ---------- API ----------
def test_generate_endpoint():
    r = client.post("/generate-quiz",
                    files=[("files", ("a.pdf", make_pdf(), "application/pdf")),
                           ("files", ("b.docx", make_docx(), "application/octet-stream"))],
                    data={"material_ids": ["m1", "m2"], "num_questions": "5", "difficulty": "easy"})
    assert r.status_code == 200, r.text
    qs = r.json()["questions"]
    assert len(qs) == 5
    assert all(len(q["options"]) == 4 and q["source_material_id"] in ("m1", "m2") for q in qs)

def test_generate_errors():
    bad = client.post("/generate-quiz", files=[("files", ("a.exe", b"x", "application/octet-stream"))],
                      data={"material_ids": ["m1"]})
    assert bad.status_code == 422
    mismatch = client.post("/generate-quiz", files=[("files", ("a.txt", LOREM.encode(), "text/plain"))],
                           data={"material_ids": ["m1", "m2"]})
    assert mismatch.status_code == 400

def test_export_report():
    body = {"quiz_title": "Quiz 1", "class_name": "CS101",
            "results": [{"student_name": "Asha", "student_email": "a@x.com", "score": 8, "total": 10,
                         "time_taken_seconds": 600, "submitted_at": "2026-10-07 10:00"},
                        {"student_name": "Ravi", "score": 3, "total": 10}],
            "question_stats": [{"question": "Q1", "correct_count": 1, "total_attempts": 2}]}
    r = client.post("/export-report", json=body)
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Results", "Summary", "Question Analysis"]
    assert wb["Results"]["E2"].value == 80.0
    assert wb["Summary"]["B4"].value == 55.0
